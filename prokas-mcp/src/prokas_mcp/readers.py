"""Read Berichtswesen exports (.xlsx/.xls/.txt/.csv) into a raw grid of cells.

Nothing here knows about columns; it only turns a file into rows plus format facts
(encoding, separator, sheet) that the parser and validate_mapping report back.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

EXCEL_EXT = {".xlsx", ".xlsm"}
XLS_EXT = {".xls"}
TEXT_EXT = {".txt", ".csv", ".tsv", ".prn", ".lst", ".dat"}
SUPPORTED_EXT = EXCEL_EXT | XLS_EXT | TEXT_EXT

SEPARATORS = ["\t", ";", "|", ","]
SEP_NAMES = {"\t": "Tab", ";": "Semikolon", "|": "Senkrechtstrich", ",": "Komma",
             "fixed": "feste Breite (Leerzeichen)"}


class ReadError(Exception):
    pass


@dataclass
class RawTable:
    rows: list[list[Any]]
    format: str
    encoding: str | None = None
    separator: str | None = None
    sheet: str | None = None
    notes: list[str] = field(default_factory=list)

    def describe(self) -> dict:
        d = {"format": self.format}
        if self.encoding:
            d["encoding"] = self.encoding
        if self.separator:
            d["separator"] = SEP_NAMES.get(self.separator, repr(self.separator))
        if self.sheet:
            d["sheet"] = self.sheet
        if self.notes:
            d["notes"] = self.notes
        return d


def read_raw(path: str | Path) -> RawTable:
    path = Path(path)
    if not path.exists():
        raise ReadError(f"Datei nicht gefunden: {path}")
    if path.stat().st_size == 0:
        raise ReadError(f"Datei ist leer: {path.name}")
    ext = path.suffix.lower()
    if ext in EXCEL_EXT:
        return _read_xlsx(path)
    if ext in XLS_EXT:
        try:
            return _read_xls(path)
        except ReadError:
            # Some tools write tab text or HTML with an .xls name; try text.
            t = _read_text(path)
            t.notes.append(".xls-Datei war keine echte Excel-Datei, als Text gelesen.")
            return t
    return _read_text(path)


def _read_xlsx(path: Path) -> RawTable:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # zipfile.BadZipFile, InvalidFileException, ...
        raise ReadError(f"Excel-Datei {path.name} kann nicht gelesen werden: {e}") from e
    try:
        best: tuple[int, str, list[list[Any]]] | None = None
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            filled = sum(1 for r in rows if any(c not in (None, "") for c in r))
            if best is None or filled > best[0]:
                best = (filled, ws.title, rows)
    finally:
        wb.close()
    if not best or best[0] == 0:
        raise ReadError(f"Excel-Datei {path.name} enthält keine Daten.")
    notes = []
    if len(wb.sheetnames) > 1:
        notes.append(f"{len(wb.sheetnames)} Blätter, gelesen: {best[1]!r}")
    return RawTable(rows=best[2], format="xlsx", sheet=best[1], notes=notes)


def _read_xls(path: Path) -> RawTable:
    try:
        import xlrd
    except ImportError as e:  # pragma: no cover - dependency is declared
        raise ReadError("Für .xls-Dateien fehlt das Paket xlrd.") from e
    try:
        book = xlrd.open_workbook(str(path))
    except Exception as e:
        raise ReadError(f"XLS-Datei {path.name} kann nicht gelesen werden: {e}") from e
    best = None
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    except Exception:
                        row.append(cell.value)
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        filled = sum(1 for r in rows if any(c not in (None, "") for c in r))
        if best is None or filled > best[0]:
            best = (filled, sh.name, rows)
    if not best or best[0] == 0:
        raise ReadError(f"XLS-Datei {path.name} enthält keine Daten.")
    return RawTable(rows=best[2], format="xls", sheet=best[1])


def decode_bytes(data: bytes) -> tuple[str, str]:
    """BOM -> UTF-8 (strict) -> Windows-1252 -> Latin-1."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace"), "utf-8-sig"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1252"), "windows-1252"
    except UnicodeDecodeError:
        return data.decode("latin-1"), "latin-1"


def detect_separator(lines: list[str]) -> str:
    """Pick the separator that splits the most lines into the same number of cells."""
    sample = [ln for ln in lines if ln.strip()][:200]
    if not sample:
        return "fixed"
    best_sep, best_score = "fixed", 0.0
    for sep in SEPARATORS:
        counts = [ln.count(sep) for ln in sample]
        nonzero = [c for c in counts if c > 0]
        if not nonzero:
            continue
        mode, freq = Counter(nonzero).most_common(1)[0]
        score = freq / len(sample) * (1 + min(mode, 20) / 20)
        if sep == ",":
            score *= 0.6  # decimal commas make comma the least trustworthy guess
        if score > best_score:
            best_sep, best_score = sep, score
    if best_score < 0.3:
        return "fixed"
    return best_sep


_FIXED_SPLIT = re.compile(r"\s{2,}")


def _read_text(path: Path) -> RawTable:
    text, enc = decode_bytes(path.read_bytes())
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x0c", "\n")
    lines = text.split("\n")
    sep = detect_separator(lines)
    rows: list[list[Any]] = []
    if sep == "fixed":
        for ln in lines:
            ln = ln.rstrip()
            rows.append([c.strip() for c in _FIXED_SPLIT.split(ln.strip())] if ln.strip() else [])
    else:
        reader = csv.reader(io.StringIO("\n".join(lines)), delimiter=sep, quotechar='"')
        for r in reader:
            rows.append([c.strip() for c in r])
    fmt = "csv" if path.suffix.lower() == ".csv" else "txt"
    return RawTable(rows=rows, format=fmt, encoding=enc, separator=sep)


def cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()
